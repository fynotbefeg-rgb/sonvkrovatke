"""План действий для Романа на 8–9 октября."""
import sys
from docx import Document
from docx.shared import Pt, RGBColor
d = Document()
st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(12)
def p(t, b=False, i=False, gray=False):
    para = d.add_paragraph(); r = para.add_run(t); r.bold = b; r.italic = i
    if gray: r.font.color.rgb = RGBColor(110, 110, 110)
def ck(t): d.add_paragraph("☐ " + t)
def code(t):
    para = d.add_paragraph(); r = para.add_run(t); r.font.name = "Courier New"; r.font.size = Pt(10)

d.add_heading("План действий: 2 рилса", 0)
p("Рилс 1 — ManyChat: автоответ в директ по слову из комментария. Рилс 2 — Claude бесплатно на год для стартапов.", i=True)

d.add_heading("Шаг 1. Согласовать сценарии", 1)
ck("Прочитать «Сценарии ManyChat и Claude» и написать правки одним сообщением (или «ок»).")
ck("Выбрать кодовое слово для ManyChat. В сценарии — «БОТ». Если меняем, скажи до генерации.")

d.add_heading("Шаг 2. Создать 2 аватара в HeyGen — делает Виктор", 1)
p("Из записей 8 октября (папка на Google Диске). Если HeyGen попросит подтверждение личности (видео-согласие), понадобится Роман — 1 минута.")
ck("Аватар «Основа» — файл dji_mimo_20261008_094328 (44 с, по пояс, с жестами).")
ck("Аватар «Хуки» — файл dji_mimo_20261008_095144 (33 с, крупный план).")
p("Голос — тот же, что сейчас (Roman, ElevenLabs). Настройки и тариф не меняем.", gray=True)

d.add_heading("Шаг 3. Сгенерировать 8 роликов", 1)
p("Тексты берём из конца документа со сценариями, раздел «Тексты для HeyGen». Вставлять как есть: английские слова уже написаны так, как их правильно прочитает озвучка.")
p("Подсказка для движений (Motion prompt) — одна на все ролики:", b=True)
code("Calm, confident, friendly delivery. Looks straight into the camera the whole time. Small natural head nods on key points, slight smile at the beginning and at the end, raises eyebrows briefly when emphasizing. Hands relaxed below the chest, occasional calm open-palm gesture at chest level. Never raises hands to the face, no pointing, no turning away.")
t = d.add_table(rows=1, cols=3); t.style = "Table Grid"
for c, x in zip(t.rows[0].cells, ["Рилс", "Ролик → имя файла", "Аватар"]): c.text = x
for r in [("ManyChat", "Хук 1 → manychat-h1", "Хуки (095144)"), ("ManyChat", "Хук 2 → manychat-h2", "Хуки (095144)"),
          ("ManyChat", "Хук 3 → manychat-h3", "Хуки (095144)"), ("ManyChat", "Основа → manychat-osnova", "Основа (094328)"),
          ("Claude на год", "Хук 1 → startups-h1", "Хуки (095144)"), ("Claude на год", "Хук 2 → startups-h2", "Хуки (095144)"),
          ("Claude на год", "Хук 3 → startups-h3", "Хуки (095144)"), ("Claude на год", "Основа → startups-osnova", "Основа (094328)")]:
    cells = t.add_row().cells
    for c, x in zip(cells, r): c.text = x
p("")
ck("Каждый ролик скачать в 1080p и загрузить в одну папку на Диске с именами из таблицы.")
p("Если основа длиннее 60 секунд и HeyGen не даёт сгенерировать её одним роликом — разбей текст на 2 части по абзацам (osnova-1, osnova-2), склею сам.", gray=True)

d.add_heading("Шаг 4. Настроить ManyChat (до публикации рилса 1)", 1)
ck("Пройти инструкцию «Настройка ManyChat» и включить запись экрана во время настройки.")
ck("Проверить со второго аккаунта: пишешь «БОТ» → приходит директ.")
ck("Прислать записи экрана и скрин с названиями кнопок.")
p("Без этого шага рилс 1 не публикуем: люди напишут «БОТ» и ничего не получат.", b=True)

d.add_heading("Шаг 5. Монтаж (делаем мы)", 1)
p("Как только ролики в папке: монтаж обоих рилсов по 3 версии (с разными хуками), без кружка с лицом, с субтитрами и вставками. Записи настройки ManyChat вставим в рилс 1 вместо карточек с шагами.")

d.add_heading("Шаг 6. Публикация", 1)
ck("Рилс 2 (Claude на год) — можно сразу после монтажа.")
ck("Рилс 1 (ManyChat) — после проверки, что автоответ на «БОТ» работает.")
ck("Через 2–3 дня сравнить охваты версий с разными хуками — лучший хук берём за основу следующих рилсов.")

d.add_heading("По желанию: перезаписать эмоции", 1)
p("Не мешает сделать эти 2 рилса, но улучшит следующие. Эмоции сейчас по 7 секунд, а для аватара нужно 30–60 секунд одним дублем:")
ck("Энергия, «сейчас покажу такое» (для хуков) — 30–60 с.")
ck("Дружелюбная улыбка (для «пиши в комментариях») — 30–60 с.")
ck("Серьёзный, «вот где проблема» — 30 с.")
p("Кадр как в 095144, та же кепка и футболка, петличка, руки опущены, паузы 1–2 секунды с закрытым ртом, взгляд в камеру.", gray=True)
d.save(sys.argv[1])
