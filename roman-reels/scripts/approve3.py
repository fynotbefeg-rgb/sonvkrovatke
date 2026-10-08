"""Документ на согласование: ManyChat + Claude для стартапов (тексты из texts.json)."""
import json, os, sys
from docx import Document
from docx.shared import Pt, RGBColor
T = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "texts.json")))
d = Document()
st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(12)
def p(t, b=False, i=False, gray=False):
    para = d.add_paragraph(); r = para.add_run(t); r.bold = b; r.italic = i
    if gray: r.font.color.rgb = RGBColor(110, 110, 110)
def li(t): d.add_paragraph(t, style="List Bullet")
# произношение для HeyGen
PH = [("Claude Console", "Клод Консоль"), ("Claude Team", "Клод Тим"), ("Claude", "Клод"), ("ManyChat", "МэниЧат"), ("Instagram", "Инстаграм"),
      ("Anthropic", "Антропик"), ("API", "эй-пи-ай"), ("AWS", "эй-дабл-ю-эс"), ("Google Cloud", "Гугл Клауд"), ("Meta", "Мета")]
def heygen(x):
    for a, b in PH: x = x.replace(a, b)
    return x

def topic(n, key, title, dur, refs, facts, note=None):
    t = T[key]
    d.add_heading(f"Рилс {n}. {title}", 1)
    p(f"Длительность: {dur}", i=True)
    d.add_heading("Хуки (3 варианта для теста)", 2)
    for k in ("h1", "h2", "h3"): p(f"Хук {k[1]}: {t[k]}")
    d.add_heading("Основной текст", 2)
    for x in t["body"].split("\n"): p(x)
    d.add_heading("Референсы — залетевшие рилсы по теме", 2)
    for l, u in refs: li(f"{l} — {u}")
    d.add_heading("На чём держатся факты", 2)
    for x in facts: li(x)
    if note: p(note, i=True, gray=True)

d.add_heading("Сценарии рилсов", 0)
p("Роман, ниже два рилса: хуки, основной текст, ссылки на референсы и откуда взяты факты. В тексте — только то, что можно проверить самому.", i=True)

topic(1, "manychat", "ManyChat: автоответ в директ по слову из комментария", "≈ 1 мин 20 с",
 [("«Хочешь фокус?» — зритель пишет слово и получает директ, 482 тыс. просмотров (@digitally.sb) — основа хука 1", "https://www.instagram.com/reel/DbRXFuLHb26/"),
  ("«Пиши „блог“ и забирай курс» — RU, 44 тыс. (@bondar_idm) — основа хука 2", "https://www.instagram.com/reel/Db34NZaMgSf/"),
  ("«В конце видео я попрошу оставить комментарий» — 131 тыс. (@victor_ofrank) — основа хука 3", "https://www.instagram.com/reel/DcRIZTTOeVo/"),
  ("RU: кейс с автоматизацией и гайдом за комментарий — 187 тыс. (@maley94)", "https://www.instagram.com/reel/Db9_o-Co7tq/"),
  ("ManyChat: обзор обновления Open Reply — 103 тыс. (@benkimball.ai)", "https://www.instagram.com/reel/DcUNlNvy6J9/")],
 ["Шаги настройки (проф. аккаунт → вход через Meta → шаблон ответа на комментарий → сообщение в директ → проверка) повторяем при настройке ManyChat на твоём аккаунте. Если названия кнопок окажутся другими, поправим текст до генерации.",
  "Список контактов в ManyChat — вкладка «Contacts».",
  "ИИ-функции в ManyChat платные — это сходится во всех источниках. Точных цен в ролике нет: сторонние сайты называют разные цифры, а официальный сайт с сервера не открылся."],
 "Важно: ролик сработает, только если ManyChat уже настроен на твоём аккаунте со словом «БОТ», иначе зрители не получат обещанное. Инструкцию для отправки по слову «БОТ» подготовлю. Справочно: Instagram в России заблокирован с 2022 года.")

topic(2, "startups", "Claude бесплатно на год для стартапов", "≈ 1 мин 25 с",
 [("Исходный рилс по теме (@itmaxxer, TikTok)", "https://www.tiktok.com/@itmaxxer/video/7693927761303325984"),
  ("Официальная страница программы", "https://claude.com/programs/startups"),
  ("CNBC о запуске программы", "https://www.cnbc.com/2026/10/06/anthropic-claude-startups-program.html")],
 ["Всё про условия перепроверено на claude.com/programs/startups 8 октября 2026: год Claude Team до 5 премиум-мест (если вы новые в Team), $1 000 на API, до $45 000 у партнёров, офис-часы с Applied AI. Подходят стартапы младше 5 лет или с инвестициями за последние 2 года, венчурные деньги не обязательны. Нужны Claude Console, почта на домене сайта и описание. Большинство решений принимают за минуты, остальные — за 2–3 рабочих дня. Кредиты сгорают через 6 месяцев и работают только через Claude Console, не через AWS или Google Cloud. При фонде из партнёрской сети — до $100 000 на API.",
  "Страны: Россия и Беларусь не входят в список поддерживаемых, Казахстан и Грузия входят (anthropic.com/supported-countries).",
  "Последний абзац «совет от тех, кто подавал» заменили на проверяемые факты (сгорание кредитов и +$100K от фонда)."])

d.add_heading("Тексты для HeyGen (с произношением)", 1)
p("Вставлять как есть: английские слова записаны так, как их правильно прочитает озвучка.", i=True, gray=True)
for key, title in (("manychat", "ManyChat"), ("startups", "Claude на год")):
    d.add_heading(title, 2)
    for k, name in (("h1", "h1"), ("h2", "h2"), ("h3", "h3"), ("body", "osnova")):
        p(f"{name}:", b=True); p(heygen(T[key][k]).replace("\n", " "))
d.save(sys.argv[1])
