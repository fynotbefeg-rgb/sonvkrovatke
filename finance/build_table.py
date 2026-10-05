"""Собирает бесплатную таблицу для самозанятых: finance/tablica-samozanyatogo.xlsx.

Запуск: python finance/build_table.py (нужен openpyxl).
Правила НПД сверены с npd.nalog.ru/faq на октябрь 2026.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("tablica-samozanyatogo.xlsx")
ROWS = 500  # строк для доходов

INK = "1F2A37"
ACCENT = "2F7D5B"
SOFT = "EAF4EE"
WARN = "FBE3DF"
MONEY = '#,##0 "₽"'
MONEY_KOP = '#,##0.00 "₽"'

h1 = Font(name="Arial", size=16, bold=True, color=INK)
h2 = Font(name="Arial", size=12, bold=True, color=INK)
body = Font(name="Arial", size=11, color=INK)
mute = Font(name="Arial", size=10, color="6B7785")
head = Font(name="Arial", size=11, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor=ACCENT)
soft_fill = PatternFill("solid", fgColor=SOFT)
input_fill = PatternFill("solid", fgColor="FFF7E0")
thin = Side(style="thin", color="D5DCE3")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# --- Старт ------------------------------------------------------------------
st = wb.active
st.title = "Старт"
st.column_dimensions["A"].width = 46
st.column_dimensions["B"].width = 22
st.column_dimensions["C"].width = 48

st["A1"] = "Таблица доходов самозанятого"
st["A1"].font = h1
st["A2"] = "Считает налог, остаток налогового бонуса и сколько осталось до лимита 2,4 млн ₽ в год."
st["A2"].font = mute

st["A4"] = "Настройки"
st["A4"].font = h2
settings = [
    ("Год, за который смотреть итоги", 2026, "Поменяйте в январе", None),
    ("Остаток налогового бонуса, ₽", 10000, "Посмотрите в «Мой налог». Бонус 10 000 ₽ даётся один раз", MONEY),
    ("Лимит дохода за год, ₽", 2400000, "Лимит по закону 422-ФЗ", MONEY),
]
for i, (label, value, hint, fmt) in enumerate(settings, start=5):
    st.cell(i, 1, label).font = body
    c = st.cell(i, 2, value)
    c.font = Font(name="Arial", size=11, bold=True, color=INK)
    c.fill = input_fill
    c.border = box
    if fmt:
        c.number_format = fmt
    st.cell(i, 3, hint).font = mute
YEAR, BONUS, LIMIT = "Старт!$B$5", "Старт!$B$6", "Старт!$B$7"

st["A9"] = "Как пользоваться"
st["A9"].font = h2
steps = [
    "1. Каждую оплату записывайте на лист «Доходы»: дату, клиента, кто платит, сумму и выбит ли чек.",
    "2. Жёлтые ячейки заполняете вы, остальное таблица считает сама.",
    "3. На листе «Итоги» — налог по месяцам, срок оплаты, невыбитые чеки и остаток до лимита.",
    "4. Три строки с пометкой «Пример» удалите, когда начнёте вести свои доходы.",
]
for i, text in enumerate(steps, start=10):
    st.cell(i, 1, text).font = body
    st.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)

st["A15"] = "Главные правила"
st["A15"].font = h2
rules = [
    ("Ставка", "4% — если платит физлицо, 6% — если платит ИП или компания"),
    ("Налоговый бонус", "10 000 ₽ один раз. Пока он не закончился, ставка 3% и 4%"),
    ("Чек", "Наличные — сразу. Безналичные — не позже 9-го числа следующего месяца"),
    ("Квитанция", "Приходит в «Мой налог» до 12-го числа следующего месяца"),
    ("Оплата налога", "До 28-го числа следующего месяца. Меньше 100 ₽ переносится на следующий месяц"),
    ("Лимит", "2,4 млн ₽ дохода за календарный год. Если превысить — право на НПД теряется"),
    ("Нельзя", "Нанимать работников по трудовому договору, перепродавать товары, работать на текущего или бывшего (до 2 лет) работодателя"),
]
for i, (k, v) in enumerate(rules, start=16):
    st.cell(i, 1, k).font = Font(name="Arial", size=11, bold=True, color=INK)
    c = st.cell(i, 2, v)
    c.font = body
    c.alignment = wrap
    st.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
    st.row_dimensions[i].height = 30

st["A24"] = (
    "Правила — на октябрь 2026, источник: npd.nalog.ru. Таблица носит справочный характер и не является "
    "налоговой консультацией: точная сумма налога — в приложении «Мой налог»."
)
st["A24"].font = mute
st["A24"].alignment = wrap
st.merge_cells("A24:C24")
st.row_dimensions[24].height = 44

# --- Доходы -----------------------------------------------------------------
inc = wb.create_sheet("Доходы")
cols = [
    ("Дата", 13), ("Клиент", 26), ("Кто платит", 20), ("Сумма", 14), ("Чек выбит", 12),
    ("Ставка", 9), ("Бонус списан", 14), ("Налог", 13), ("Год", 7), ("Месяц", 7),
]
for j, (name, width) in enumerate(cols, start=1):
    c = inc.cell(1, j, name)
    c.font, c.fill, c.border = head, head_fill, box
    c.alignment = Alignment(horizontal="center", vertical="center")
    inc.column_dimensions[c.column_letter].width = width
inc.freeze_panes = "A2"

examples = [
    ("2026-10-02", "Пример: заказ на сайте", "Физлицо", 1199, "Да"),
    ("2026-10-05", "Пример: ООО «Ромашка»", "ИП или компания", 15000, "Да"),
    ("2026-10-11", "Пример: перевод на карту", "Физлицо", 2500, "Нет"),
]
last = ROWS + 1
for r in range(2, last + 1):
    if r - 2 < len(examples):
        d, client, payer, amount, check = examples[r - 2]
        from datetime import date

        inc.cell(r, 1, date.fromisoformat(d))
        inc.cell(r, 2, client)
        inc.cell(r, 3, payer)
        inc.cell(r, 4, amount)
        inc.cell(r, 5, check)
    phys = f'C{r}="Физлицо"'
    inc.cell(r, 6, f'=IF(D{r}="","",IF({phys},0.04,0.06))')
    inc.cell(r, 7, f'=IF(D{r}="","",MIN(MAX({BONUS}-SUM(G$1:G{r - 1}),0),D{r}*IF({phys},0.01,0.02)))')
    inc.cell(r, 8, f'=IF(D{r}="","",D{r}*F{r}-G{r})')
    inc.cell(r, 9, f'=IF(A{r}="","",YEAR(A{r}))')
    inc.cell(r, 10, f'=IF(A{r}="","",MONTH(A{r}))')
    for j in range(1, 11):
        c = inc.cell(r, j)
        c.font, c.border = body, box
        if j <= 5:
            c.fill = input_fill
    inc.cell(r, 1).number_format = "DD.MM.YYYY"
    inc.cell(r, 4).number_format = MONEY
    inc.cell(r, 6).number_format = "0%"
    inc.cell(r, 7).number_format = MONEY_KOP
    inc.cell(r, 8).number_format = MONEY_KOP

payer = DataValidation(type="list", formula1='"Физлицо,ИП или компания"', allow_blank=True)
check = DataValidation(type="list", formula1='"Да,Нет"', allow_blank=True)
inc.add_data_validation(payer)
inc.add_data_validation(check)
payer.add(f"C2:C{last}")
check.add(f"E2:E{last}")
inc.conditional_formatting.add(
    f"A2:H{last}",
    FormulaRule(formula=['$E2="Нет"'], fill=PatternFill("solid", fgColor=WARN)),
)

# --- Итоги ------------------------------------------------------------------
sm = wb.create_sheet("Итоги")
for col, width in zip("ABCDEF", (28, 18, 16, 16, 16, 30)):
    sm.column_dimensions[col].width = width
sm["A1"] = "Итоги за год"
sm["A1"].font = h1
sm["B1"] = f"={YEAR}"
sm["B1"].font = h1

D, H, E = f"Доходы!$D$2:$D${last}", f"Доходы!$H$2:$H${last}", f"Доходы!$E$2:$E${last}"
Y, M = f"Доходы!$I$2:$I${last}", f"Доходы!$J$2:$J${last}"
summary = [
    ("Доход за год", f"=SUMIFS({D},{Y},{YEAR})", MONEY),
    ("Налог за год", f"=SUMIFS({H},{Y},{YEAR})", MONEY_KOP),
    ("Остаток до лимита", f"={LIMIT}-B3", MONEY),
    ("Лимит использован", f"=B3/{LIMIT}", "0.0%"),
    ("Остаток налогового бонуса", f"=MAX({BONUS}-SUM(Доходы!$G$2:$G${last}),0)", MONEY_KOP),
    (
        "Статус",
        f'=IF(B6>=1,"Лимит превышен: право на НПД потеряно, нужен другой режим",'
        f'IF(B6>=0.8,"Больше 80% лимита — планируйте заранее","В пределах лимита"))',
        None,
    ),
]
for i, (label, formula, fmt) in enumerate(summary, start=3):
    sm.cell(i, 1, label).font = body
    c = sm.cell(i, 2, formula)
    c.font = Font(name="Arial", size=12, bold=True, color=INK)
    c.fill = soft_fill
    c.border = box
    if fmt:
        c.number_format = fmt
sm.merge_cells("B8:F8")
sm.conditional_formatting.add("B8", FormulaRule(formula=["$B$6>=0.8"], fill=PatternFill("solid", fgColor=WARN)))

hdr = ["Месяц", "Доход", "Налог", "Оплатить до", "Чеков не выбито", "Что сделать"]
for j, name in enumerate(hdr, start=1):
    c = sm.cell(10, j, name)
    c.font, c.fill, c.border = head, head_fill, box
months = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
          "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
for m, name in enumerate(months, start=1):
    r = 10 + m
    sm.cell(r, 1, name)
    sm.cell(r, 2, f"=SUMIFS({D},{Y},{YEAR},{M},{m})").number_format = MONEY
    sm.cell(r, 3, f"=SUMIFS({H},{Y},{YEAR},{M},{m})").number_format = MONEY_KOP
    sm.cell(r, 4, f"=DATE({YEAR},{m + 1},28)").number_format = "DD.MM.YYYY"
    sm.cell(r, 5, f'=COUNTIFS({Y},{YEAR},{M},{m},{E},"Нет")')
    sm.cell(r, 6, f'=IF(E{r}>0,"Выбить чеки до 9-го",IF(AND(C{r}>0,C{r}<100),"Меньше 100 ₽ — перенесётся",""))')
    for j in range(1, 7):
        sm.cell(r, j).font, sm.cell(r, j).border = body, box
sm.cell(23, 1, "Итого").font = Font(name="Arial", size=11, bold=True, color=INK)
sm.cell(23, 2, "=SUM(B11:B22)").number_format = MONEY
sm.cell(23, 3, "=SUM(C11:C22)").number_format = MONEY_KOP
sm.cell(23, 5, "=SUM(E11:E22)")
sm["A25"] = "Налог — оценка. Точная сумма — в квитанции «Мой налог», она приходит до 12-го числа."
sm["A25"].font = mute

wb.active = wb.sheetnames.index("Доходы")
wb.save(OUT)
print(OUT)
